from pathlib import Path

from docx import Document

from services.docx_loader import docx_to_text


def test_docx_to_text_extracts_paragraphs(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Welcome to the call.")
    doc.add_paragraph("Q: What about demand?")
    doc.add_paragraph("A: Strong.")
    target = tmp_path / "fixture.docx"
    doc.save(target)

    text = docx_to_text(target)
    assert "Welcome to the call." in text
    assert "Q: What about demand?" in text
    assert "A: Strong." in text
    assert text.count("\n") >= 2


def test_docx_to_text_rejects_non_docx(tmp_path: Path):
    target = tmp_path / "bad.txt"
    target.write_text("not docx", encoding="utf-8")
    try:
        docx_to_text(target)
    except ValueError:
        return
    raise AssertionError("expected ValueError for non-docx")
