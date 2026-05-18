import io

from docx import Document
from fastapi.testclient import TestClient

from app import app


def _make_docx_bytes() -> bytes:
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Welcome.")
    doc.add_paragraph("Q: demand?")
    doc.add_paragraph("A: strong.")
    doc.save(buf)
    return buf.getvalue()


def test_upload_starts_job(monkeypatch):
    async def fake_process(**kwargs):
        kwargs["store"].update(kwargs["job_id"],
                               status=__import__("services.job_store", fromlist=["JobStatus"]).JobStatus.DONE,
                               step="done")

    monkeypatch.setattr("routes.upload._process_upload", fake_process)
    client = TestClient(app)
    response = client.post(
        "/upload",
        files={"full_docx": ("a.docx", _make_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"case_name": "myco"},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["slug"] == "myco"
    assert payload["job_id"]
